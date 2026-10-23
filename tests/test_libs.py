"""Unit tests for TestHive's Python libraries and Studio's API (standard-library unittest, no installs).
Run by scripts/selftest.sh:  python -m unittest discover -s tests
Each test works in its own temporary project folder with a copy of scripts/; the real project is never touched."""
import os, sys, json, shutil, tempfile, threading, subprocess, unittest, urllib.request, urllib.error

for k in ("TH_STATE", "TH_ANALYSIS", "TH_CASES", "TH_JOB", "TH_ARCHIVE", "TH_CLAUDE_FAKE"): os.environ.pop(k, None)  # defaults only
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "scripts", "lib"))
import analysis, design, intake, modules, studio  # noqa: E402

BASH = os.environ.get("TH_BASH") or shutil.which("bash")  # selftest passes the Git Bash path on Windows (not WSL's bash)
REQ_HEAD = ",".join(analysis.REQ_COLS) + "\n"


class Project(unittest.TestCase):
    """A temporary project as the current folder."""

    def setUp(self):
        self.home, self.dir = os.getcwd(), tempfile.mkdtemp()
        shutil.copytree(os.path.join(ROOT, "scripts"), os.path.join(self.dir, "scripts"), ignore=shutil.ignore_patterns("__pycache__"))
        os.chdir(self.dir)

    def tearDown(self):
        os.chdir(self.home); shutil.rmtree(self.dir, ignore_errors=True)

    def stage(self, *a, env=None):
        return subprocess.run([BASH, "scripts/stage.sh", *a], capture_output=True, text=True, encoding="utf-8", env=dict(os.environ, **(env or {})))

    def write(self, path, text):
        os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
        with open(path, "w", encoding="utf-8", newline="") as f: f.write(text)


class DocumentNames(unittest.TestCase):
    def test_allowed(self):
        for n in ("prd.pdf", "specs/api.yaml", "قصص المستخدم.docx", "notes (v2).txt"): self.assertEqual(intake.doc_name(n), n)

    def test_refused(self):
        for n in ("../x.pdf", "a/../b.pdf", "/etc/passwd", "C:\\x.pdf", "..\\..\\.env", "_old/x.pdf", "_OLD/x.pdf", ".env", "con.txt", "a//b.pdf", ""):
            self.assertIsNone(intake.doc_name(n), n)


class Documents(Project):
    def test_upload_replace_remove_restore(self):
        ok, _ = intake.save_upload("SIGNUP", "prd.txt", b"version one", by="Sara"); self.assertTrue(ok)
        ok, msg = intake.save_upload("SIGNUP", "prd.txt", b"version two", by="Sara"); self.assertIn("earlier version is kept", msg)
        self.assertEqual([v["name"] for v in intake.versions("SIGNUP")], ["prd.txt"])
        ok, _ = intake.remove_source("SIGNUP", "prd.txt", "Sara"); self.assertTrue(ok)
        self.assertIsNone(intake.source_file("SIGNUP", "prd.txt"))
        removed = [v for v in intake.versions("SIGNUP") if not v["current"]]
        self.assertEqual(len(removed), 2)  # both versions are kept, neither is current
        newest = intake.versions("SIGNUP")[0]
        ok, _ = intake.restore_source("SIGNUP", newest["version"], "prd.txt", "Sara"); self.assertTrue(ok)
        with open(intake.source_file("SIGNUP", "prd.txt"), "rb") as f: self.assertEqual(f.read(), b"version two")
        self.assertEqual([h["Action"].split()[0] for h in intake.history("SIGNUP")], ["uploaded", "replaced", "removed", "restored"])
        self.assertTrue(intake.old_file("SIGNUP", newest["version"], "prd.txt"), "the restored version stays in _old/")

    def test_history_is_safe_to_open(self):
        intake.save_upload("SIGNUP", "a.txt", b"x", by="=HYPERLINK(1)")
        with open(os.path.join(intake.SOURCES, "SIGNUP", "_old", "history.csv"), "a", encoding="utf-8") as f: f.write("not a date,removed,a.txt,x\n")
        self.assertEqual([h["By"] for h in intake.history("SIGNUP")], ["'=HYPERLINK(1)"], "no formula, and the damaged line is skipped")
        intake.sources("SIGNUP", "2026-01-01T00:00:00Z")  # a damaged history never stops the module from opening

    def test_nothing_outside_the_module(self):
        self.write(".env", "SECRET=1"); intake.save_upload("SIGNUP", "a.txt", b"x")
        for n in ("../../.env", "..\\..\\.env", "a.txt/../../../.env"):
            self.assertIsNone(intake.source_file("SIGNUP", n)); self.assertFalse(intake.remove_source("SIGNUP", n, "x")[0])
        self.assertFalse(intake.restore_source("SIGNUP", "../../x", "a.txt", "x")[0])
        self.assertFalse(intake.save_upload("SIGNUP", "run.exe", b"MZ")[0])

    def test_reading_forgets_removed_documents(self):
        intake.save_upload("SIGNUP", "a.txt", b"Rule A: the password has 8 characters.")
        intake.save_upload("SIGNUP", "b.txt", b"Rule B: the code has 6 digits.")
        env = dict(os.environ, TH_BASH=BASH)
        run = lambda: subprocess.run([BASH, "scripts/ingest.sh", "SIGNUP"], capture_output=True, text=True, env=env)
        self.assertEqual(run().returncode, 0)
        self.assertIn("6 digits", intake.source_text("SIGNUP", "b.txt")["text"])
        intake.remove_source("SIGNUP", "b.txt", "Sara")
        out = os.path.join("analysis", "sources-text", "SIGNUP")
        self.assertFalse(any("b.txt" in f for f in os.listdir(out)), "the removed document's text is gone at once")
        os.remove(os.path.join("requirements", "sources", "SIGNUP", "a.txt")); intake.save_upload("SIGNUP", "c.txt", b"Rule C.")
        self.assertEqual(run().returncode, 0)
        self.assertEqual(sorted(f for f in os.listdir(out) if f.endswith(".md")), ["c.txt.md"], "reading again drops text of files that are gone")


class ReviewPacks(Project):
    """An Excel workbook sent to reviewers may come back late, or in two copies: only what each reviewer changed is taken."""

    def setUp(self):
        super().setUp(); self.stage("init", "LOGIN")
        self.write(analysis.GAP_FILE, ",".join(analysis.GAP_COLS) + "\n" +
                   "".join(f"GAP-LOGIN-00{i},LOGIN,,Missing,High,Question {i}?,,,Open\n" for i in (1, 2, 3)))

    def pack(self, name, edits):
        import openpyxl
        analysis.cmd_workbook(["LOGIN", name]); wb = openpyxl.load_workbook(name); ws = wb["Gaps"]
        head = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2):
            for col, v in edits.get(row[0].value, {}).items(): row[head.index(col)].value = v
        wb.save(name)

    def answers(self): return {g["ID"]: g["Answer"] for g in analysis.read_csv(analysis.GAP_FILE, analysis.GAP_COLS)}

    def studio(self, gid, text):
        self.assertFalse(analysis.apply_answers("LOGIN", [("s", {"ID": gid, "Answer": text, "Answered By": "Sara", "Status": "Answered"})])[2])

    def test_late_packs_keep_newer_answers(self):
        self.pack("a.xlsx", {"GAP-LOGIN-001": {"Answer": "A1", "Answered By": "Ann"}})
        self.pack("b.xlsx", {"GAP-LOGIN-002": {"Answer": "B2", "Answered By": "Bob"}})
        self.studio("GAP-LOGIN-003", "S3")  # answered in Studio after both packs went out
        analysis.cmd_import(["LOGIN", "a.xlsx"]); analysis.cmd_import(["LOGIN", "b.xlsx"])
        self.assertEqual(self.answers(), {"GAP-LOGIN-001": "A1", "GAP-LOGIN-002": "B2", "GAP-LOGIN-003": "S3"})

    def test_the_same_answer_changed_twice_is_refused(self):
        self.pack("c.xlsx", {"GAP-LOGIN-001": {"Answer": "C1", "Answered By": "Cam"}})
        self.studio("GAP-LOGIN-001", "S1")
        import io, contextlib
        out = io.StringIO()
        with self.assertRaises(SystemExit), contextlib.redirect_stdout(out): analysis.cmd_import(["LOGIN", "c.xlsx"])
        self.assertIn("also in TestHive", out.getvalue()); self.assertEqual(self.answers()["GAP-LOGIN-001"], "S1")


class Archive(Project):
    def setUp(self):
        super().setUp()
        for m in ("SIGNUP", "LOGIN"): self.stage("init", m)
        self.write(analysis.REQ_FILE, REQ_HEAD + "REQ-SIGNUP-001,SIGNUP,prd,1,a,Functional,P0,1,Draft,\nREQ-LOGIN-001,LOGIN,prd,1,b,Functional,P0,1,Draft,\n")
        intake.save_upload("SIGNUP", "prd.txt", b"doc"); self.write("reviews/DESIGN-SIGNUP-1.md", "review")

    def test_archive_and_restore(self):
        r = self.stage("archive", "SIGNUP", "--by", "Sara", "--comment", "added by mistake"); self.assertEqual(r.returncode, 0, r.stdout)
        self.assertNotIn("SIGNUP", studio.load_state()["modules"])
        self.assertEqual([x["ID"] for x in analysis.read_csv(analysis.REQ_FILE, analysis.REQ_COLS)], ["REQ-LOGIN-001"])
        self.assertFalse(os.path.exists("requirements/sources/SIGNUP") or os.path.exists("reviews/DESIGN-SIGNUP-1.md"))
        [a] = modules.listing(); self.assertEqual((a["module"], a["by"], a["rows"]), ("SIGNUP", "Sara", 1))
        r = self.stage("restore", a["name"], "--by", "Sara"); self.assertEqual(r.returncode, 0, r.stdout)
        self.assertIn("SIGNUP", studio.load_state()["modules"])
        self.assertEqual(len(analysis.read_csv(analysis.REQ_FILE, analysis.REQ_COLS)), 2)
        self.assertTrue(os.path.exists("requirements/sources/SIGNUP/prd.txt") and os.path.exists("reviews/DESIGN-SIGNUP-1.md"))
        self.assertEqual(modules.listing(), [])
        with open("workflow/APPROVALS.md", encoding="utf-8") as f: log = f.read()
        self.assertIn("| Archived | Sara | added by mistake", log); self.assertIn("| Restored | Sara |", log)

    def test_restore_refused_when_ids_exist_again(self):
        self.stage("archive", "SIGNUP", "--by", "Sara", "--comment", "x")
        with open(analysis.REQ_FILE, "a", encoding="utf-8") as f: f.write("REQ-SIGNUP-001,OTHER,prd,1,c,Functional,P0,1,Draft,\n")
        r = self.stage("restore", modules.listing()[0]["name"], "--by", "Sara")
        self.assertNotEqual(r.returncode, 0); self.assertIn("already has REQ-SIGNUP-001", r.stdout)
        self.assertTrue(modules.listing(), "a refused restore leaves the archive as it was")

    def test_agents_cannot_archive(self):
        r = self.stage("archive", "SIGNUP", "--by", "Bot", "--comment", "x", env={"TH_JOB": "job-1"})
        self.assertEqual(r.returncode, 2); self.assertIn("SIGNUP", studio.load_state()["modules"])

    def test_archive_needs_a_reason(self):
        self.assertNotEqual(self.stage("archive", "SIGNUP", "--by", "Sara").returncode, 0)

    def test_changed_or_planted_archive_is_refused(self):
        self.stage("archive", "SIGNUP", "--by", "Sara", "--comment", "x")
        [a] = modules.listing(); dest = os.path.join("archive", a["name"])
        with open(os.path.join(dest, "state.json"), encoding="utf-8") as f: st = json.load(f)
        st["stages"]["design"]["status"] = "approved"  # an approval nobody gave
        self.write(os.path.join(dest, "state.json"), json.dumps(st))
        r = self.stage("restore", a["name"], "--by", "Sara")
        self.assertNotEqual(r.returncode, 0); self.assertIn("fingerprint", r.stdout)
        planted = os.path.join("archive", "GHOST-20260101-000000")
        self.write(os.path.join(planted, "ARCHIVE.json"), json.dumps({"module": "GHOST", "archived_at": "x", "files": [], "rows": {}}))
        self.write(os.path.join(planted, "state.json"), json.dumps(st))
        self.assertIn("fingerprint", self.stage("restore", "GHOST-20260101-000000", "--by", "Sara").stdout)

    def test_only_the_modules_own_paths(self):
        for p in ("requirements/sources/SIGNUP", "reviews/DESIGN-SIGNUP-2.md", "runs/agents/20261022-101010-SIGNUP-manual-tester"):
            self.assertTrue(modules.belongs(p, "SIGNUP"), p)
        for p in ("../.env", "/etc/passwd", "C:/x", "requirements/sources/LOGIN", "reviews/DESIGN-LOGIN-1.md", "workflow/state.json",
                  "requirements/sources/SIGNUP/../../../.env", "runs/agents/20261022-101010-SIGNUP-x/../../../.env"):
            self.assertFalse(modules.belongs(p, "SIGNUP"), p)

    def test_a_locked_file_undoes_the_archive(self):
        real, calls = shutil.move, []
        def move(src, dst):
            calls.append(src)
            if len(calls) == 2: raise PermissionError(13, "in use", src)
            return real(src, dst)
        modules.shutil.move = move
        try:
            with self.assertRaises(SystemExit): modules.archive("SIGNUP", os.path.join("archive", "SIGNUP-20260101-000000"), "Sara", "x")
        finally: modules.shutil.move = real
        self.assertTrue(os.path.exists("requirements/sources/SIGNUP/prd.txt") and os.path.exists("reviews/DESIGN-SIGNUP-1.md"))
        self.assertEqual(len(analysis.read_csv(analysis.REQ_FILE, analysis.REQ_COLS)), 2, "the registers are untouched")
        self.assertFalse(os.path.exists("archive/SIGNUP-20260101-000000"))


class ApprovedDesign(Project):
    def test_review_refused_after_approval(self):
        self.stage("init", "SIGNUP"); os.makedirs(design.TC_DIR)
        with open(studio.STATE, encoding="utf-8") as f: s = json.load(f)
        s["modules"]["SIGNUP"]["stages"]["design"]["status"] = "approved"
        self.write(studio.STATE, json.dumps(s))
        errors = design.apply_review("SIGNUP", [("case", {"ID": "TC-SIGNUP-001", "Review": "Remove"})], [])[-1]
        self.assertTrue(errors and "read-only" in errors[0])
        s["modules"]["SIGNUP"]["stages"]["design"]["status"] = "in_progress"; self.write(studio.STATE, json.dumps(s))
        self.assertEqual(design.apply_review("SIGNUP", [], [("request", {"Request": "more cases", "Requested By": "Sara"})])[-1], [])


class StudioApi(Project):
    """Studio's HTTP API on a free port, as the page uses it."""

    def setUp(self):
        super().setUp()
        os.environ["TH_BASH"] = BASH
        os.environ["TH_STUDIO_CONFIG"] = os.path.join(self.dir, "studio-config.json")  # never the user's own settings
        studio.PROJECT["path"] = self.dir
        self.srv = studio.ThreadingHTTPServer(("127.0.0.1", 0), studio.Handler)
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()
        self.base = f"http://127.0.0.1:{self.srv.server_address[1]}"

    def tearDown(self):
        self.srv.shutdown(); self.srv.server_close(); super().tearDown()

    def call(self, path, body=None, token=True, raw=None, ctype="application/json"):
        data = raw if raw is not None else (None if body is None else json.dumps(body).encode())
        req = urllib.request.Request(self.base + path, data=data, method="GET" if data is None else "POST",
                                     headers={**({"X-TestHive-Token": studio.TOKEN} if token else {}), **({"Content-Type": ctype} if data is not None else {})})
        try:
            with urllib.request.urlopen(req) as r: return r.status, r.read(), r.headers
        except urllib.error.HTTPError as e:
            with e: return e.code, e.read(), e.headers

    def js(self, *a, **k):
        code, data, _ = self.call(*a, **k); return code, json.loads(data or b"{}")

    def test_page_and_guards(self):
        code, _, h = self.call("/", token=False)
        self.assertEqual(code, 200); self.assertIn("script-src 'self';", h["Content-Security-Policy"]); self.assertNotIn("unsafe-inline'; img", h["Content-Security-Policy"].split("script-src")[1])
        self.assertEqual(self.call("/app.js", token=False)[0], 200)
        for path, ctype in (("/style.css", "text/css"), ("/js/core.js", "text/javascript"), ("/fonts/ibm-plex-sans-400.woff2", "font/woff2")):
            code, _, h = self.call(path, token=False); self.assertEqual((code, h["Content-Type"].split(";")[0]), (200, ctype), path)
        self.assertEqual(self.call("/js/nothing.js", token=False)[0], 404)
        for path in ("/fonts/OFL.txt", "/js/../index.html", "/..%2fscripts/lib/studio.py", "/JS/core.js"):  # only page files, never anything else
            self.assertEqual(self.call(path, token=False)[0], 403, path)
        self.assertEqual(self.call("/api/overview", token=False)[0], 403)
        self.assertEqual(self.call("/api/nothing")[0], 404)
        self.assertEqual(self.call("/api/module?m=NOPE")[0], 404)
        self.assertEqual(self.call("/api/project", raw=b"x=1", ctype="text/plain")[0], 415)
        self.assertEqual(self.call("/api/project", raw=b"[1]")[0], 400)
        self.assertTrue(all(fn.__doc__ is None or fn.__doc__ for fn, _ in studio.ROUTES.values()))

    def test_module_documents_and_archive(self):
        self.assertEqual(self.js("/api/module-new", {"module": "signup", "title": "Sign-up"})[0], 200)
        code, r = self.js("/api/source?m=SIGNUP&name=prd.txt&by=Sara", raw=b"Rule: 8 characters.", ctype="application/octet-stream")
        self.assertEqual(code, 200, r)
        code, data, h = self.call("/api/source-file?m=SIGNUP&name=prd.txt")
        self.assertEqual((code, data), (200, b"Rule: 8 characters.")); self.assertEqual(h["Content-Security-Policy"], "sandbox")
        self.assertEqual(self.call("/api/source-file?m=SIGNUP&name=..%2F..%2Fworkflow%2Fstate.json")[0], 404)
        self.assertEqual(self.js("/api/source-remove", {"module": "SIGNUP", "name": "prd.txt"})[0], 400, "a removal is signed")
        self.assertEqual(self.js("/api/source-remove", {"module": "SIGNUP", "name": "prd.txt", "by": "Sara"})[0], 200)
        [v] = self.js("/api/source-versions?m=SIGNUP")[1]["versions"]
        self.assertEqual(self.js("/api/source-restore", {"module": "SIGNUP", "version": v["version"], "name": "prd.txt", "by": "Sara"})[0], 200)
        self.assertEqual(self.js("/api/module-archive", {"module": "SIGNUP", "by": "Sara", "comment": "x", "confirm": "LOGIN"})[0], 400)
        self.assertEqual(self.js("/api/module-archive", {"module": "SIGNUP", "by": "Sara", "confirm": "SIGNUP"})[0], 400, "a reason is needed")
        code, r = self.js("/api/module-archive", {"module": "SIGNUP", "by": "Sara", "comment": "duplicate", "confirm": "signup"})
        self.assertEqual(code, 200, r)
        self.assertEqual(self.js("/api/overview")[1]["modules"], [])
        [a] = self.js("/api/archive")[1]["archived"]
        self.assertEqual(self.js("/api/module-restore", {"archive": "../../x", "by": "Sara"})[0], 400)
        self.assertEqual(self.js("/api/module-restore", {"archive": a["name"], "by": "Sara"})[0], 200)
        self.assertEqual([m["module"] for m in self.js("/api/overview")[1]["modules"]], ["SIGNUP"])


class StudioProjects(StudioApi):
    """One Studio program for any project: until a project folder is open only the project routes answer."""

    def test_open_a_project(self):
        studio.PROJECT["path"] = None
        code, r = self.js("/api/overview"); self.assertEqual((code, r.get("no_project")), (409, True))
        self.assertEqual(self.js("/api/open-project", {"path": os.path.join(self.dir, "nothing")})[0], 400)
        other = os.path.join(self.dir, "other"); os.makedirs(os.path.join(other, "templates")); shutil.copytree("scripts", os.path.join(other, "scripts"))
        code, r = self.js("/api/open-project", {"path": other}); self.assertEqual(code, 200, r)
        self.assertEqual(os.path.realpath(os.getcwd()), os.path.realpath(other))
        self.assertEqual(self.js("/api/overview")[0], 200)
        self.assertIn(os.path.realpath(other), self.js("/api/projects")[1]["recent"])
        self.assertEqual(studio.jobs.JOBS, {})  # the runs of the project left behind are not shown in this one
        os.chdir(self.dir)


class StudioWithoutBash(StudioApi):
    """Review mode on a PC with Python only: project setup, a module, documents, reading them and the stage rules need no bash."""

    def setUp(self):
        super().setUp()
        shutil.copytree(os.path.join(ROOT, "templates"), "templates")
        os.environ["TH_BASH"] = os.path.join(self.dir, "no-bash-here")

    def tearDown(self):
        os.environ["TH_BASH"] = BASH; super().tearDown()

    def test_setup_to_stage_rules(self):
        code, r = self.js("/api/project", {"project": "Demo", "app": "Demo", "platforms": ["Web"]})
        self.assertEqual(code, 200, r); self.assertTrue(os.path.exists("PROJECT_BRIEF.md"))
        self.assertEqual(self.js("/api/module-new", {"module": "LOGIN", "title": "Log in"})[0], 200)
        self.assertEqual(self.js("/api/source?m=LOGIN&name=prd.txt&by=Sara", raw=b"Rule: 8 characters.", ctype="application/octet-stream")[0], 200)
        code, r = self.js("/api/ingest", {"module": "LOGIN", "mask": False}); self.assertEqual(code, 200, r)
        code, r = self.js("/api/stage", {"module": "LOGIN", "stage": "analysis", "action": "approve", "by": "Sara", "comment": ""})
        self.assertEqual(code, 400); self.assertIn("not waiting for approval", r.get("message", "") + r.get("error", ""))

    def test_review_mode(self):
        fake = os.environ.pop("TH_CLAUDE_FAKE", None)
        try:
            o = self.js("/api/overview")[1]
            self.assertEqual(o["agents"], "review"); self.assertIn("Git Bash", o["agents_missing"])
            self.js("/api/project", {"project": "Demo", "app": "Demo", "platforms": ["Web"]}); self.js("/api/module-new", {"module": "LOGIN"})
            code, r = self.js("/api/run", {"module": "LOGIN", "action": "analyse"})
            self.assertEqual(code, 409); self.assertIn("Review mode", r["message"])
        finally:
            if fake: os.environ["TH_CLAUDE_FAKE"] = fake


class JobsOfAnotherStudio(Project):
    """A second Studio on the same project sees a running job of the first one and does not mark it interrupted."""

    def job(self, pid):
        studio.jobs.JOBS.clear()
        studio.jobs.save({"id": "J1", "module": "LOGIN", "kind": "agent", "status": "running", "started": "x", "pid": pid})
        studio.jobs.load_history(); return studio.jobs.JOBS["J1"]["status"]

    def test_live_job_stays_running(self):
        p = subprocess.Popen([sys.executable, "-c", "import time; time.sleep(30)"])
        try:
            self.assertEqual(self.job(p.pid), "running"); self.assertTrue(studio.jobs.running("LOGIN"))
            self.assertTrue(studio.jobs.stop("J1")); p.wait(10)  # Stop reaches a job this Studio did not start
        finally: p.kill(); studio.jobs.JOBS.clear()

    def test_dead_job_is_interrupted(self):
        p = subprocess.Popen([sys.executable, "-c", "pass"]); p.wait()
        self.assertEqual(self.job(p.pid), "interrupted"); studio.jobs.JOBS.clear()


class StageLock(Project):
    def test_lock_left_by_a_crash_is_taken_over(self):
        self.assertEqual(self.stage("init", "LOGIN").returncode, 0)
        self.write("workflow/state.json.lock", ""); os.utime("workflow/state.json.lock", (1, 1))
        r = self.stage("status", "LOGIN"); self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertFalse(os.path.exists("workflow/state.json.lock"))


class StudioExplore(StudioApi):
    """The built product: what Explore found is shown, Explore needs the web tools, and a module with only an address is analysed from it."""

    def module(self):
        shutil.copytree(os.path.join(ROOT, "templates"), "templates")
        self.js("/api/project", {"project": "Demo", "app": "Demo", "platforms": ["Web"], "base_url": "https://test.example.com"})
        self.js("/api/module-new", {"module": "LOGIN", "url": "https://test.example.com/login"})

    def test_explore_result_and_report(self):
        self.module(); d = self.js("/api/module?m=LOGIN")[1]
        self.assertIsNone(d["exploration"]); self.assertEqual(d["project_url"], "https://test.example.com")
        self.assertEqual(d["next"]["title"], "Analyse the built product")  # no documents, only an address
        self.write("analysis/exploration/LOGIN/web/exploration.json", json.dumps([{"url": "https://test.example.com/login", "n": "01"}, {"url": "x", "error": "Timeout"}]))
        self.write("analysis/exploration/LOGIN/web/exploration.md", "# Web exploration")
        e = self.js("/api/module?m=LOGIN")[1]["exploration"]
        self.assertEqual((e["pages"], e["failed"], e["url"]), (1, 1, "https://test.example.com/login"))
        self.assertEqual(self.js("/api/exploration?m=LOGIN")[1]["text"], "# Web exploration")
        self.assertFalse(self.js("/api/module?m=LOGIN")[1]["test_login"], "the template's change-me is not a login")
        with open(".env", "a", encoding="utf-8") as f: f.write("TEST_USER=tester\nTEST_USER_PASSWORD=Secret-123\n")
        d = self.js("/api/module?m=LOGIN")[1]; self.assertTrue(d["test_login"]); self.assertNotIn("Secret-123", json.dumps(d))

    def test_explore_needs_the_web_tools(self):
        self.module(); self.assertFalse(self.js("/api/module?m=LOGIN")[1]["web_tools"])  # a temporary project has no web/node_modules
        code, r = self.js("/api/run", {"module": "LOGIN", "action": "explore-web"})
        self.assertEqual(code, 409); self.assertIn("web tools", r["message"])


@unittest.skipUnless(os.path.isdir(os.path.join(ROOT, "web", "node_modules", "@playwright", "test")), "web tools not installed")
class ExploreLogin(unittest.TestCase):
    """Explore logs in with the test account from .env, never follows logout, and never writes the password down."""

    def explore(self, password):
        out = tempfile.mkdtemp(); site = os.path.join(ROOT, "tests", "fixtures", "login-site", "login.html")
        url = "file:///" + site.replace(os.sep, "/").lstrip("/")
        env = dict(os.environ, TH_LOGIN_USER="tester", TH_LOGIN_PASSWORD=password)
        subprocess.run(["node", os.path.join(ROOT, "web", "explore.mjs"), url, out], capture_output=True, text=True, env=env, cwd=os.path.join(ROOT, "web"), timeout=180)
        with open(os.path.join(out, "exploration.md"), encoding="utf-8") as f: md = f.read()
        shutil.rmtree(out, ignore_errors=True); return md

    def test_logs_in_and_explores_behind_the_login(self):
        md = self.explore("Secret-123")
        self.assertIn("Login: logged in with TEST_USER", md); self.assertIn("Welcome to the dashboard", md); self.assertIn("Settings page", md)
        self.assertNotIn("You are logged out", md, "the logout link is not followed"); self.assertNotIn("Secret-123", md)

    def test_a_wrong_password_is_reported(self):
        md = self.explore("wrong-one")
        self.assertIn("Login: failed", md); self.assertNotIn("Welcome to the dashboard", md); self.assertNotIn("wrong-one", md)


class StudioPort(unittest.TestCase):
    def test_a_second_studio_cannot_share_the_port(self):  # it takes a free one instead (Windows allowed sharing before)
        a = studio.Server(("127.0.0.1", 0), studio.Handler)
        try:
            with self.assertRaises(OSError): studio.Server(("127.0.0.1", a.server_address[1]), studio.Handler)
        finally: a.server_close()


class StudioProgramRoute(StudioWithoutBash):
    """The same as StudioWithoutBash, but every library starts the others the way the single Studio program does (--lib)."""

    def setUp(self):
        super().setUp(); os.environ["TH_LAUNCHER"] = os.path.join(ROOT, "scripts", "lib", "studio_app.py")

    def tearDown(self):
        os.environ.pop("TH_LAUNCHER", None); super().tearDown()

    def test_archive_and_restore(self):  # stage.py runs modules.py: a library started by a library
        self.js("/api/project", {"project": "Demo", "app": "Demo", "platforms": ["Web"]}); self.js("/api/module-new", {"module": "LOGIN"})
        code, r = self.js("/api/module-archive", {"module": "LOGIN", "by": "Sara", "comment": "test", "confirm": "LOGIN"})
        self.assertEqual(code, 200, r)
        arch = self.js("/api/archive")[1]["archived"]
        code, r = self.js("/api/module-restore", {"archive": arch[0]["name"], "by": "Sara"})
        self.assertEqual(code, 200, r); self.assertIn("LOGIN", self.js("/api/overview")[1]["modules"][0]["module"])

    def test_unknown_library_refused(self):
        r = subprocess.run([sys.executable, os.environ["TH_LAUNCHER"], "--lib", "../x.py"], capture_output=True, text=True)
        self.assertNotEqual(r.returncode, 0); self.assertIn("Unknown library", r.stderr)


if __name__ == "__main__":
    unittest.main()
