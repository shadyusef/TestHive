"""TestHive ingest: turns requirement sources into plain text the analyst can read. Called by scripts/ingest.sh."""
import sys, os, re, json, csv, glob, time

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

TEXT = (".txt", ".md")
IMAGES = (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp")
OUT = os.environ.get("TH_ANALYSIS", "analysis")


def die(m): print(m); sys.exit(1)


def read_pdf(p):
    from pypdf import PdfReader
    parts = []
    for i, page in enumerate(PdfReader(p).pages, 1):
        t = (page.extract_text() or "").strip()
        if t: parts.append(f"<!-- page {i} -->\n{t}")
    return "\n\n".join(parts)


def read_docx(p):
    import docx
    d, parts = docx.Document(p), []
    body = d.element.body
    # keep document order: paragraphs and tables interleaved
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = docx.text.paragraph.Paragraph(child, d)
            t = para.text.strip()
            if not t: continue
            style = (para.style.name or "").lower() if para.style is not None else ""
            m = re.match(r"heading (\d)", style)
            parts.append(("#" * int(m.group(1)) + " " + t) if m else t)
        elif tag == "tbl":
            table = docx.table.Table(child, d)
            rows = [[c.text.strip().replace("\n", " ") for c in r.cells] for r in table.rows]
            if rows:
                parts.append("| " + " | ".join(rows[0]) + " |\n|" + "---|" * len(rows[0]))
                parts += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(parts)


def read_xlsx(p):
    import openpyxl
    wb, parts = openpyxl.load_workbook(p, read_only=True, data_only=True), []
    for ws in wb.worksheets:
        parts.append(f"## Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if v is None else str(v).replace("\n", " ") for v in row]
            if any(cells): parts.append("| " + " | ".join(cells) + " |")
    return "\n".join(parts)


def read_csv(p):
    with open(p, encoding="utf-8-sig", newline="") as f:
        return "\n".join("| " + " | ".join(r) + " |" for r in csv.reader(f))


def load_structured(p):
    with open(p, encoding="utf-8-sig") as f: raw = f.read()
    if p.lower().endswith(".json"): return json.loads(raw)
    import yaml
    return yaml.safe_load(raw)


def resolve(spec, obj):
    seen = 0
    while isinstance(obj, dict) and "$ref" in obj and seen < 20:
        ref = obj["$ref"]
        if not ref.startswith("#/"): return obj
        node = spec
        for part in ref[2:].split("/"): node = node.get(part, {}) if isinstance(node, dict) else {}
        obj, seen = node, seen + 1
    return obj


def schema_lines(spec, schema, indent="  ", depth=0):
    schema = resolve(spec, schema or {})
    if depth > 4 or not isinstance(schema, dict): return []
    out, req = [], set(schema.get("required", []))
    for name, prop in (schema.get("properties") or {}).items():
        prop = resolve(spec, prop)
        limits = [f"{k}={prop[k]}" for k in ("type", "format", "minLength", "maxLength", "minimum", "maximum", "pattern", "enum", "default") if k in prop]
        out.append(f"{indent}- `{name}`{' (required)' if name in req else ''}: {', '.join(map(str, limits))}" + (f" — {prop['description']}" if prop.get("description") else ""))
        if prop.get("type") == "object" or "properties" in prop: out += schema_lines(spec, prop, indent + "  ", depth + 1)
    return out


def read_openapi(spec):
    info = spec.get("info", {})
    parts = [f"# API: {info.get('title', '')} {info.get('version', '')}".strip(), info.get("description", "") or ""]
    for path, ops in (spec.get("paths") or {}).items():
        for method, op in ops.items():
            if method.lower() not in ("get", "post", "put", "patch", "delete"): continue
            parts.append(f"\n## {method.upper()} {path}" + (f" — {op.get('summary')}" if op.get("summary") else ""))
            if op.get("description"): parts.append(op["description"])
            for prm in op.get("parameters", []):
                prm = resolve(spec, prm)
                parts.append(f"- parameter `{prm.get('name')}` in {prm.get('in')}{' (required)' if prm.get('required') else ''}: {prm.get('schema', {})}")
            body = (((op.get("requestBody") or {}).get("content") or {}).get("application/json") or {}).get("schema")
            if body: parts += ["Request body:"] + schema_lines(spec, body)
            for code, resp in (op.get("responses") or {}).items():
                resp = resolve(spec, resp)
                parts.append(f"- response {code}: {resp.get('description', '')}")
                rs = (((resp.get("content") or {}).get("application/json") or {}).get("schema"))
                if rs: parts += schema_lines(spec, rs, "    ")
    return "\n".join(p for p in parts if p is not None)


def convert(path):
    """Returns (kind, text). text is None for images (read directly by Claude)."""
    ext = os.path.splitext(path)[1].lower()
    if ext in TEXT:
        with open(path, encoding="utf-8-sig") as f: return "text", f.read()
    if ext == ".pdf": return "pdf", read_pdf(path)
    if ext == ".docx": return "word", read_docx(path)
    if ext in (".xlsx", ".xlsm"): return "excel", read_xlsx(path)
    if ext == ".csv": return "csv", read_csv(path)
    if ext in (".json", ".yaml", ".yml"):
        spec = load_structured(path)
        if isinstance(spec, dict) and ("openapi" in spec or "swagger" in spec): return "openapi", read_openapi(spec)
        return "data", json.dumps(spec, indent=2, ensure_ascii=False)
    if ext in IMAGES: return "image", None
    if ext == ".doc": die(f"{path}: old Word format (.doc). Open it in Word and save as .docx.")
    return None, None


def main():
    a = sys.argv[1:]
    mask = "--mask" in a
    a = [x for x in a if x != "--mask"]
    if len(a) != 1: die("Usage: ingest.sh <MODULE> [--mask]   (reads requirements/sources/<module>/ or requirements/sources/)")
    module = a[0].upper()
    src_dir = next((d for d in (f"requirements/sources/{a[0]}", f"requirements/sources/{module}", f"requirements/sources/{a[0].lower()}") if os.path.isdir(d)), "requirements/sources")
    files = sorted(p for p in glob.glob(os.path.join(src_dir, "**", "*"), recursive=True)
                   if os.path.isfile(p) and not os.path.basename(p).startswith(".") and ".masked." not in p
                   and "_old" not in p.replace(os.sep, "/").split("/"))  # _old/: earlier versions kept by Studio
    if not files: die(f"No source files in {src_dir}/. Put the module's documents there first.")
    out_dir = os.path.join(OUT, "sources-text", module); os.makedirs(out_dir, exist_ok=True)
    if mask: import mask as masker
    rows, skipped = [], []
    for p in files:
        try: kind, text = convert(p)
        except Exception as e: skipped.append(f"{p}: could not read ({e})"); continue
        if kind is None: skipped.append(f"{p}: unknown file type"); continue
        rel = p.replace(os.sep, "/")
        base = re.sub(r"[^\w.-]+", "_", os.path.basename(p))
        if kind == "image":
            out_file, n = "", 0
            if mask:
                out_file = os.path.join(out_dir, base.rsplit(".", 1)[0] + ".masked.png")
                masker.mask_image(p, out_file)
            rows.append([rel, kind, out_file.replace(os.sep, "/") if out_file else rel, "", "yes" if mask else "no"]); continue
        header = f"<!-- source: {rel} | type: {kind} | ingested: {time.strftime('%Y-%m-%d')} -->\n"
        out_file = os.path.join(out_dir, base + ".md")
        if mask: text = masker.mask_string(text)[0]
        with open(out_file, "w", encoding="utf-8", newline="\n") as f: f.write(header + text.strip() + "\n")
        rows.append([rel, kind, out_file.replace(os.sep, "/"), len(text.split()), "yes" if mask else "no"])
    made = {os.path.realpath(r[2]) for r in rows}
    for old in glob.glob(os.path.join(out_dir, "*.md")) + glob.glob(os.path.join(out_dir, "*.masked.png")):
        if os.path.realpath(old) not in made: os.remove(old)  # text of a document that is gone: the analyst must not read it
    idx = os.path.join(out_dir, "index.csv")
    with open(idx, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n"); w.writerow(["Source", "Type", "Read from", "Words", "Masked"]); w.writerows(rows)
    print(f"INGESTED {len(rows)} sources for {module} -> {out_dir.replace(os.sep, '/')}/ (index.csv lists them)")
    for r in rows: print(f"  {r[1]:<8} {r[0]}" + ("  (image: Claude reads it directly)" if r[1] == "image" else f"  ({r[3]} words)"))
    for s in skipped: print("  SKIPPED " + s)
    sys.exit(1 if skipped and not rows else 0)


if __name__ == "__main__":
    main()
