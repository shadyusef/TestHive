"""TestHive test design: check the case sheet against the approved analysis baseline, set risk-based priority,
Excel review workbook and import, change impact, score against a condition key, design baseline.
Called by scripts/check-design.sh, design-workbook.sh, import-review.sh, design-impact.sh, score-design.sh and stage.sh.
Formats: docs/FORMATS.md (test case sheet, design requests, design workbook, design baselines)."""
import sys, os, re, csv, json, time
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from analysis import read_csv, write_csv, refs, die, stage_status, add_sheet, dropdown, REQ_COLS, GAP_COLS, DIR as ADIR

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
TC_DIR = os.environ.get("TH_CASES", "test-cases")
CASE_FILE, ASK_FILE = os.path.join(TC_DIR, "TestCases.csv"), os.path.join(TC_DIR, "requests.csv")
CASE_COLS = ["ID", "Module", "Platform", "Title", "Priority", "Type", "Requirements", "Technique", "Risk", "Tags",
             "Preconditions", "Steps", "Expected Result", "Test Data", "Assumption", "Origin", "Review", "Review Comment",
             "Automation", "Automation Spec", "Status"]
ASK_COLS = ["ID", "Module", "Requirement", "Request", "Requested By", "Status", "Response"]
PLATFORMS = ["Android", "iOS", "Both", "API", "Web"]
TYPES = ["Positive", "Negative", "Localization", "Accessibility", "Security", "Performance"]  # "Functional" (older sheets) counts as Positive
TECHNIQUES = ["Equivalence", "Boundary", "Decision table", "State transition", "Error guessing", "Use case", "Checklist"]
RISKS = ["High", "Medium", "Low"]
TAGS = ["smoke", "sanity", "regression", "e2e"]
ORIGINS = ["Designed", "Human", "Requested"]
REVIEWS = ["Accepted", "Change", "Remove", "Updated"]
STATUSES = ["Draft", "Ready", "Needs update", "Removed", "Automated", "Done"]
ASK_STATUS = ["Open", "Done"]
EDITABLE = ["Title", "Priority", "Type", "Requirements", "Technique", "Risk", "Tags", "Platform", "Preconditions", "Steps",
            "Expected Result", "Test Data", "Assumption"]  # content the human may change directly in Excel
# a requirement with a numeric limit needs at least one Boundary case
LIMIT = re.compile(r"\b(at least|at most|minimum|maximum|min|max|minlength|maxlength|or older|older than|up to)\b|"
                   r"\d+[\s-]*(characters?|chars?|digits?|seconds?|minutes?|hours?|days?|times|attempts?)\b", re.I)
VAGUE = re.compile(r"\b(works?|correct(ly)?|properly|as expected|appropriate(ly)?|successfully|fine|normal(ly)?)\b", re.I)


def module_rows(rows, module):
    return [r for r in rows if r["Module"].upper() == module or r["ID"].startswith(f"TC-{module}-")]


def baseline(module, version=None):
    """(version, requirements, gaps) of the approved analysis baseline: the latest one, or the given vN."""
    base = os.path.join(ADIR, "baseline", module)
    vs = sorted((int(d[1:]) for d in os.listdir(base) if re.fullmatch(r"v\d+", d)), reverse=True) if os.path.isdir(base) else []
    if version: vs = [int(version.lstrip("v"))] if int(version.lstrip("v")) in vs else []
    if not vs: return None
    d = os.path.join(base, f"v{vs[0]}")
    return f"v{vs[0]}", read_csv(os.path.join(d, "requirements.csv"), REQ_COLS), read_csv(os.path.join(d, "gaps.csv"), GAP_COLS)


def all_gaps(module, base_gaps):
    """The baseline's gaps, plus questions raised later (e.g. "From design: ...") in the live register for the next analysis round."""
    live = {g["ID"]: g for g in read_csv(os.path.join(ADIR, "gaps.csv"), GAP_COLS) if g["Module"].upper() == module}
    live.update({g["ID"]: g for g in base_gaps})
    return list(live.values())


def suggested_priority(case, req_by_id):
    """Risk-based priority: impact (highest priority of the linked requirements) x likelihood (the case's Risk)."""
    impact = max(({"P0": 3, "P1": 2}.get(req_by_id[x]["Priority"], 1) for x in refs(case["Requirements"]) if x in req_by_id), default=0)
    likely = {"High": 3, "Medium": 2, "Low": 1}.get(case["Risk"], 0)
    if not impact or not likely: return None
    s = impact * likely
    return "P0" if s == 9 else "P1" if s == 6 else "P2" if s >= 3 else "P3"  # the usual 3x3 risk matrix


def active(c): return c["Status"] != "Removed"


def is_negative(c): return c["Type"] == "Negative"


def coverage(reqs, cases):
    """Per active requirement: the active cases linked to it."""
    return {r["ID"]: [c for c in cases if active(c) and r["ID"] in refs(c["Requirements"])] for r in reqs if r["Status"] != "Removed"}


def needs_negative(r): return r["Priority"] == "P0" and r["Type"] in ("Functional", "Business rule", "API")


def needs_boundary(r): return bool(LIMIT.search(r["Text"]))


def problems(module, for_approval=False):
    b = baseline(module)
    if b is None: return [f"no approved analysis baseline for {module}: approve the analysis stage first ({ADIR}/baseline/{module}/)"], [], None, [], []
    ver, reqs, gaps = b
    req_by_id = {r["ID"]: r for r in reqs}
    gaps = all_gaps(module, gaps)
    gap_by_id = {g["ID"]: g for g in gaps}
    all_cases = read_csv(CASE_FILE, CASE_COLS)
    cases = module_rows(all_cases, module)
    asks = [q for q in read_csv(ASK_FILE, ASK_COLS) if q["Module"].upper() == module]
    err, warn, seen, closed = [], [], set(), {}
    for c in cases:
        i = c["ID"] or "?"
        if not re.fullmatch(rf"TC-{module}(-[A-Z0-9]+)*-\d{{3}}", i): err.append(f"{i}: ID must look like TC-{module}-001")
        if i in seen: err.append(f"{i}: duplicate ID")
        seen.add(i)
        for col, values in (("Platform", PLATFORMS), ("Type", TYPES + ["Functional"]), ("Technique", TECHNIQUES), ("Risk", RISKS),
                            ("Origin", ORIGINS), ("Status", STATUSES)):
            if not c[col]: err.append(f"{i}: empty {col}")
            elif c[col] not in values: err.append(f"{i}: {col} must be one of {', '.join(values)}")
        for col in ("Title", "Steps", "Expected Result"):
            if not c[col]: err.append(f"{i}: empty {col}")
        if not c["Priority"]: err.append(f"{i}: empty Priority (run check-design {module} --set-priority)")
        elif not re.fullmatch(r"P[0-3]", c["Priority"]): err.append(f"{i}: Priority must be P0-P3")
        tags = refs(c["Tags"])
        if not tags: err.append(f"{i}: empty Tags (at least one of {', '.join(TAGS)})")
        bad_tags = [t for t in tags if t not in TAGS]
        if bad_tags: err.append(f"{i}: unknown tag {', '.join(bad_tags)} (allowed: {', '.join(TAGS)})")
        if c["Review"] and c["Review"] not in REVIEWS: err.append(f"{i}: Review must be empty or one of {', '.join(REVIEWS)}")
        if not active(c): continue
        if not refs(c["Requirements"]):
            (err if for_approval else warn).append(f"{i}: linked to no requirement (orphan case)")
        for x in refs(c["Requirements"]):
            if x not in req_by_id: err.append(f"{i}: points to {x}, which is not in analysis baseline {ver}")
            elif req_by_id[x]["Status"] == "Removed": err.append(f"{i}: points to {x}, which is Removed in baseline {ver}: remove or relink the case")
        for x in refs(c["Assumption"]):
            if x not in gap_by_id: err.append(f"{i}: assumption {x} is not a gap of {module} (baseline {ver} or analysis/gaps.csv)")
            elif gap_by_id[x]["Status"] == "Closed": closed.setdefault(x, []).append(i)
        if VAGUE.search(c["Expected Result"]): warn.append(f"{i}: vague expected result ('{VAGUE.search(c['Expected Result']).group(0)}'): state exactly what is seen")
        sp = suggested_priority(c, req_by_id)
        if sp and c["Priority"] and sp != c["Priority"]: warn.append(f"{i}: priority {c['Priority']} differs from the risk rule ({sp})")
        if "smoke" in tags and c["Priority"] not in ("P0", "P1"): warn.append(f"{i}: smoke case with priority {c['Priority']}")
        if for_approval:
            if c["Status"] == "Needs update": err.append(f"{i}: Needs update (a requirement changed); update the case, then set Status Draft")
            if c["Review"] in ("Change", "Remove"): err.append(f"{i}: the reviewer asked to {c['Review'].lower()} it; the designer has not done it yet")
    if closed: warn.append(f"{sum(map(len, closed.values()))} cases cite answered (Closed) gaps as assumptions; use the answers instead: "
                           + ", ".join(f"{g} ({len(c)} cases)" for g, c in sorted(closed.items())))
    cov = coverage(reqs, cases)
    for r in reqs:
        if r["Status"] == "Removed": continue
        linked = cov[r["ID"]]
        if not linked: err.append(f"{r['ID']}: no test case ({r['Priority']}, {r['Type']}: {r['Text'][:70]})"); continue
        if needs_negative(r) and not any(is_negative(c) for c in linked): err.append(f"{r['ID']}: P0 {r['Type'].lower()} with no Negative case")
        if needs_boundary(r) and not any(c["Technique"] == "Boundary" for c in linked):
            (err if r["Priority"] == "P0" else warn).append(f"{r['ID']}: has a numeric limit but no Boundary case")
    used = {x for c in cases if active(c) for x in refs(c["Assumption"])}
    loose = [g["ID"] for g in gaps if g["Status"] == "Open" and g["ID"] not in used and any(req_by_id.get(x, {}).get("Status") != "Removed" for x in refs(g["Requirement"]))]
    if loose: warn.append(f"{len(loose)} open gaps are not recorded as an Assumption on any case: {', '.join(loose)}")
    aseen = set()
    for q in asks:
        i = q["ID"] or "?"
        if not re.fullmatch(rf"ASK-{module}-\d{{3}}", i): err.append(f"{i}: request ID must look like ASK-{module}-001")
        if i in aseen: err.append(f"{i}: duplicate request ID")
        aseen.add(i)
        if q["Status"] not in ASK_STATUS: err.append(f"{i}: request Status must be Open or Done")
        if q["Status"] == "Done" and not q["Response"]: err.append(f"{i}: Done but no Response (name the cases added or say why not)")
        if for_approval and q["Status"] == "Open": err.append(f"{i}: request still Open: {q['Request'][:80]}")
    if for_approval and not [c for c in cases if active(c)]: err.append(f"no active test cases for {module}")
    return err, warn, ver, reqs, cases


def cmd_check(a):
    module = next((x.upper() for x in a if not x.startswith("--")), None)
    if not module: die("Usage: check-design.sh <MODULE> [--set-priority] [--for-approval]")
    if "--set-priority" in a and baseline(module):
        req_by_id = {r["ID"]: r for r in baseline(module)[1]}
        rows, n = read_csv(CASE_FILE, CASE_COLS), 0
        for c in module_rows(rows, module):
            if not c["Priority"] and (sp := suggested_priority(c, req_by_id)): c["Priority"] = sp; n += 1
        if n: write_csv(CASE_FILE, CASE_COLS, rows)
        print(f"priority set by the risk rule on {n} cases")
    err, warn, ver, reqs, cases = problems(module, "--for-approval" in a)
    for w in warn: print("  warning: " + w)
    if err: print(f"DESIGN NOT OK (module {module})\n" + "\n".join("- " + e for e in err)); sys.exit(1)
    act = [c for c in cases if active(c)]
    cov = coverage(reqs, cases)
    count = lambda col, values: ", ".join(f"{v} {n}" for v in values if (n := sum(1 for c in act if c[col] == v)))
    print(f"DESIGN OK (module {module}, analysis baseline {ver}): {len(cases)} cases ({len(act)} active), "
          f"{sum(1 for v in cov.values() if v)} of {len(cov)} requirements covered")
    print(f"  by type: {count('Type', TYPES + ['Functional'])}")
    print(f"  by technique: {count('Technique', TECHNIQUES)}")
    print(f"  by priority: {count('Priority', ['P0', 'P1', 'P2', 'P3'])}; by origin: {count('Origin', ORIGINS)}")
    print(f"  tags: " + ", ".join(f"{t} {sum(1 for c in act if t in refs(c['Tags']))}" for t in TAGS))


def cmd_workbook(a):
    import openpyxl
    from openpyxl.styles import Font
    if not a: die("Usage: design-workbook.sh <MODULE> [out.xlsx]")
    module = a[0].upper()
    out = a[1] if len(a) > 1 else os.path.join(TC_DIR, f"{module}-design.xlsx")
    b = baseline(module)
    if b is None: die(f"No approved analysis baseline for {module}.")
    ver, reqs, gaps = b
    gaps = all_gaps(module, gaps)
    cases = module_rows(read_csv(CASE_FILE, CASE_COLS), module)
    if not cases: die(f"No test cases for {module} in {CASE_FILE}")
    asks = [q for q in read_csv(ASK_FILE, ASK_COLS) if q["Module"].upper() == module]
    cov = coverage(reqs, cases)
    act = [c for c in cases if active(c)]
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Cover"
    covered = sum(1 for v in cov.values() if v)
    for row in [("TestHive test design workbook", ""), ("Module", module), ("Created", time.strftime("%Y-%m-%d %H:%M")),
                ("Analysis baseline", ver), ("Test cases", f"{len(act)} active ({len(cases) - len(act)} removed)"),
                ("Requirements covered", f"{covered} of {len(cov)} ({100 * covered // max(len(cov), 1)}%)"),
                ("By priority", ", ".join(f"{p} {sum(1 for c in act if c['Priority'] == p)}" for p in ("P0", "P1", "P2", "P3"))),
                ("Smoke / sanity / regression / e2e", " / ".join(str(sum(1 for c in act if t in refs(c["Tags"]))) for t in TAGS)), ("", ""),
                ("How to review", "Sheet 'Cases': for each case choose Review = Accepted, Change (write what to change in Review Comment) or Remove."),
                ("", "You may also edit a case's text directly (Steps, Expected Result, ...): TestHive marks it Updated."),
                ("", "To add a case yourself: add a row at the bottom with an empty ID; fill at least Title, Steps and Expected Result."),
                ("", "To ask TestHive for more cases: sheet 'Requests', a row with an empty ID; fill Request and Requested By."),
                ("", "Save the file, then run: import-review " + module + ". Do not change IDs or sheet names."),
                ("", "'Coverage' shows which cases test each requirement; 'Assumptions' shows open questions the cases rely on.")]:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=14); ws.column_dimensions["A"].width = 30; ws.column_dimensions["B"].width = 110
    ccols = ["ID", "Requirements", "Title", "Priority", "Type", "Technique", "Risk", "Tags", "Platform", "Preconditions", "Steps",
             "Expected Result", "Test Data", "Assumption", "Origin", "Status", "Review", "Review Comment"]
    add_sheet(wb.create_sheet("Cases"), ccols, [[c[x] for x in ccols] for c in cases],
              [15, 17, 40, 8, 11, 14, 8, 14, 9, 28, 50, 45, 25, 15, 10, 11, 11, 35], editable=("Review", "Review Comment"))
    cs = wb["Cases"]
    for h, values in (("Review", REVIEWS[:3]), ("Priority", ["P0", "P1", "P2", "P3"]), ("Type", TYPES), ("Technique", TECHNIQUES),
                      ("Risk", RISKS), ("Platform", PLATFORMS)):
        dropdown(cs, openpyxl.utils.get_column_letter(ccols.index(h) + 1), values)
    gap_by_req = {}
    for g in gaps:
        if g["Status"] == "Open":
            for x in refs(g["Requirement"]): gap_by_req.setdefault(x, []).append(g["ID"])
    add_sheet(wb.create_sheet("Coverage"), ["Requirement", "Priority", "Type", "Text", "Cases", "Case IDs", "Negative", "Boundary", "Open gaps"],
              [[r["ID"], r["Priority"], r["Type"], r["Text"], len(cov[r["ID"]]), ", ".join(c["ID"] for c in cov[r["ID"]]),
                ("yes" if any(is_negative(c) for c in cov[r["ID"]]) else "NO") if needs_negative(r) else "-",
                ("yes" if any(c["Technique"] == "Boundary" for c in cov[r["ID"]]) else "NO") if needs_boundary(r) else "-",
                ", ".join(gap_by_req.get(r["ID"], []))] for r in reqs if r["Status"] != "Removed"],
              [16, 8, 13, 60, 7, 40, 9, 9, 18])
    add_sheet(wb.create_sheet("Assumptions"), ["Gap", "Severity", "Requirement", "Question", "Cases relying on it"],
              [[g["ID"], g["Severity"], g["Requirement"], g["Question"],
                ", ".join(c["ID"] for c in act if g["ID"] in refs(c["Assumption"]))] for g in gaps if g["Status"] == "Open"],
              [16, 9, 18, 70, 40])
    add_sheet(wb.create_sheet("Requests"), ASK_COLS[:1] + ASK_COLS[2:], [[q[x] for x in ASK_COLS if x != "Module"] for q in asks],
              [14, 18, 60, 18, 9, 50], editable=("Requirement", "Request", "Requested By"))
    add_sheet(wb.create_sheet("Requirements"), ["ID", "Priority", "Type", "Status", "Text"],
              [[r["ID"], r["Priority"], r["Type"], r["Status"], r["Text"]] for r in reqs], [16, 8, 13, 11, 90])
    os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    wb.save(out)
    print(f"WORKBOOK {out.replace(os.sep, '/')}: {len(act)} active cases, {covered} of {len(cov)} requirements covered, {len(asks)} requests")


def apply_review(module, case_rows, ask_rows, source="Excel"):
    """Applies the human's review to TestCases.csv and requests.csv, all or nothing. case_rows / ask_rows: [(label, {column: value})].
    A case row: ID with Review / Review Comment (and any EDITABLE column the human changed), or an empty ID for a new case.
    A request row: an empty ID with Request and Requested By. Used by the Excel import and by TestHive Studio.
    Returns (decided, edited, new_cases, new_asks, errors); nothing is written if errors."""
    if stage_status(module, "design") == "approved":  # the approved cases are a frozen baseline: a change starts with a reopen
        return {r: 0 for r in REVIEWS}, 0, 0, 0, [f"The design of {module} is approved, so its cases are read-only. "
                                                  "To change them, reopen the design (Approvals) and say why."]
    rows = read_csv(CASE_FILE, CASE_COLS)
    by_id = {c["ID"]: c for c in rows}
    nums = [int(c["ID"].rsplit("-", 1)[1]) for c in rows if re.fullmatch(rf"TC-{module}-\d{{3}}", c["ID"])]
    allowed = {"Priority": ["P0", "P1", "P2", "P3"], "Type": TYPES + ["Functional"], "Technique": TECHNIQUES, "Risk": RISKS, "Platform": PLATFORMS}
    errors, decided, edited, new_cases = [], {r: 0 for r in REVIEWS}, 0, 0
    for n, v in case_rows:
        v = {h: str(x or "").strip() for h, x in v.items()}
        cid, review, note = v.get("ID", ""), v.get("Review", ""), v.get("Review Comment", "")
        bad = [f"{h} '{v[h]}'" for h, vals in allowed.items() if v.get(h) and v[h] not in vals]
        bad += [f"tag '{t}'" for t in refs(v.get("Tags", "")) if t not in TAGS]
        if bad: errors.append(f"{n} ({cid or 'new'}): not allowed: {', '.join(bad)}"); continue
        unchanged = cid in by_id and review == by_id[cid]["Review"]  # e.g. "Updated", written by the designer
        if review and review not in REVIEWS[:3] and not unchanged: errors.append(f"{n} ({cid or 'new'}): Review must be Accepted, Change or Remove"); continue
        if review == "Change" and not note: errors.append(f"{n} ({cid}): Change needs a Review Comment saying what to change"); continue
        if not cid:
            if not v.get("Title"): continue
            if not (v.get("Steps") and v.get("Expected Result")): errors.append(f"{n}: a new case needs Title, Steps and Expected Result"); continue
            nums.append(max(nums, default=0) + 1)
            c = {k: "" for k in CASE_COLS}
            c.update({h: v.get(h, "") for h in EDITABLE})
            c.update({"ID": f"TC-{module}-{nums[-1]:03d}", "Module": module, "Origin": "Human", "Status": "Draft", "Automation": "Candidate",
                      "Review": "Accepted", "Review Comment": note})
            rows.append(c); new_cases += 1; continue
        c = by_id.get(cid)
        if c is None or not module_rows([c], module): errors.append(f"{n}: {cid} is not a case of {module}"); continue
        changes = [h for h in EDITABLE if h in v and v[h] != c[h]]
        if changes:  # the human's own edit; a Change or Remove decision on the same row still counts
            for h in changes: c[h] = v[h]
            edited += 1
            c["Review"] = review if review in ("Change", "Remove") else "Updated"
            if review in ("Change", "Remove"): decided[review] += 1
            c["Review Comment"] = (note + " " if note else "") + f"[edited in {source}: {', '.join(changes)}]"
            continue
        if review and (review, note) != (c["Review"], c["Review Comment"]):
            c["Review"], c["Review Comment"] = review, note; decided[review] += 1
        elif note != c["Review Comment"]: c["Review Comment"] = note
    asks = read_csv(ASK_FILE, ASK_COLS)
    anums = [int(q["ID"].rsplit("-", 1)[1]) for q in asks if re.fullmatch(rf"ASK-{module}-\d{{3}}", q["ID"])]
    new_asks = 0
    for n, v in ask_rows:
        v = {h: str(x or "").strip() for h, x in v.items()}
        if v.get("ID") or not v.get("Request"): continue
        if not v.get("Requested By"): errors.append(f"{n}: Requested By is empty"); continue
        anums.append(max(anums, default=0) + 1)
        asks.append({"ID": f"ASK-{module}-{anums[-1]:03d}", "Module": module, "Requirement": v.get("Requirement", ""), "Request": v["Request"],
                     "Requested By": v["Requested By"], "Status": "Open", "Response": ""}); new_asks += 1
    if not errors:
        write_csv(CASE_FILE, CASE_COLS, rows)
        if new_asks: write_csv(ASK_FILE, ASK_COLS, asks)
    return decided, edited, new_cases, new_asks, errors


def cmd_import(a):
    import openpyxl
    if not a: die("Usage: import-review.sh <MODULE> [workbook.xlsx]")
    module = a[0].upper()
    path = a[1] if len(a) > 1 else os.path.join(TC_DIR, f"{module}-design.xlsx")
    if not os.path.exists(path): die(f"Workbook not found: {path}. Create it with design-workbook {module}.")
    wb = openpyxl.load_workbook(path, data_only=True)
    for s in ("Cases", "Requests"):
        if s not in wb.sheetnames: die(f"The workbook has no '{s}' sheet. Do not rename sheets.")

    def table(name, need, label):
        rows = list(wb[name].iter_rows(values_only=True))
        head = [str(h or "").strip() for h in rows[0]]
        if any(h not in head for h in need): die(f"The '{name}' sheet is missing a column: " + ", ".join(h for h in need if h not in head))
        return [(f"{label} {n}", {h: ("" if row[i] is None else str(row[i]).strip()) for i, h in enumerate(head) if h})
                for n, row in enumerate(rows[1:], 2) if any(row)]

    decided, edited, new_cases, new_asks, errors = apply_review(
        module, table("Cases", ["ID", "Title", "Steps", "Expected Result", "Review", "Review Comment"], "row"),
        table("Requests", ["ID", "Request", "Requested By"], "Requests row"))
    if errors: die("IMPORT REFUSED, nothing changed:\n" + "\n".join("- " + e for e in errors))
    print(f"IMPORTED {path.replace(os.sep, '/')}: {decided['Accepted']} accepted, {decided['Change']} to change, {decided['Remove']} to remove, "
          f"{edited} edited in Excel, {new_cases} new cases, {new_asks} new requests. Next: the designer re-checks.")


def cmd_impact(a):
    module = next((x.upper() for x in a if not x.startswith("--") and not re.fullmatch(r"v\d+", x)), None)
    if not module: die("Usage: design-impact.sh <MODULE> [--from vN] [--to vN] [--apply]")
    opt = {a[i][2:]: a[i + 1] for i in range(len(a) - 1) if a[i] in ("--from", "--to")}
    meta_dir = os.path.join(TC_DIR, "baseline", module)
    dvs = sorted((int(d[1:]) for d in os.listdir(meta_dir) if re.fullmatch(r"v\d+", d)), reverse=True) if os.path.isdir(meta_dir) else []
    frm = opt.get("from")
    if not frm and dvs:
        with open(os.path.join(meta_dir, f"v{dvs[0]}", "meta.json"), encoding="utf-8") as f: frm = json.load(f).get("analysis_baseline")
    if not frm: die(f"No design baseline for {module} yet, so the old analysis version is unknown: give --from vN")
    old, new = baseline(module, frm), baseline(module, opt.get("to"))
    if not old or not new: die(f"Analysis baseline {frm if not old else opt.get('to', 'latest')} not found for {module}")
    o, nw = {r["ID"]: r for r in old[1]}, {r["ID"]: r for r in new[1]}
    removed = [i for i in o if o[i]["Status"] != "Removed" and (i not in nw or nw[i]["Status"] == "Removed")]
    changed = [i for i in o if i in nw and i not in removed and o[i]["Status"] != "Removed" and
               any(o[i][k] != nw[i][k] for k in ("Text", "Version", "Priority", "Type"))]
    added = [i for i in nw if nw[i]["Status"] != "Removed" and (i not in o or o[i]["Status"] == "Removed")]
    rows = read_csv(CASE_FILE, CASE_COLS)
    hit = [c for c in module_rows(rows, module) if active(c) and set(refs(c["Requirements"])) & set(removed + changed)]
    print(f"IMPACT {module}: analysis {old[0]} -> {new[0]}: {len(changed)} changed, {len(removed)} removed, {len(added)} new requirements; {len(hit)} cases affected")
    for i in changed: print(f"  changed {i} (v{o[i]['Version']} -> v{nw[i]['Version']}): {nw[i]['Text'][:110]}")
    for i in removed: print(f"  removed {i}: {o[i]['Text'][:110]}")
    for i in added: print(f"  new     {i}: {nw[i]['Text'][:110]} (needs cases)")
    for c in hit: print(f"  case {c['ID']} ({', '.join(x for x in refs(c['Requirements']) if x in removed + changed)}): {c['Title'][:90]}")
    if "--apply" in a and hit:
        for c in hit:
            why = ", ".join(x for x in refs(c["Requirements"]) if x in removed + changed)
            c["Status"] = "Needs update"
            c["Review Comment"] = (c["Review Comment"] + " " if c["Review Comment"] else "") + f"[impact: {why} changed in analysis {new[0]}]"
        write_csv(CASE_FILE, CASE_COLS, rows)
        print(f"  {len(hit)} cases set to Needs update. Next: the designer updates them (and adds cases for new requirements).")


def cmd_score(a):
    if len(a) < 2: die("Usage: score-design.sh <MODULE> <condition-key.csv>")
    module, key_path = a[0].upper(), a[1]
    cases = [c for c in module_rows(read_csv(CASE_FILE, CASE_COLS), module) if active(c)]
    with open(key_path, encoding="utf-8-sig", newline="") as f: keys = list(csv.DictReader(f))

    def has(t, alt):  # starts a word; a number must also end there ("7" does not match "17" or "70")
        alt = alt.strip().lower()
        return re.search(r"(?<![a-z0-9])" + re.escape(alt) + (r"(?![0-9])" if alt[-1:].isdigit() else ""), t) is not None
    # what the case checks is in its title and expected result; steps and test data also fill in other fields, which would match too much
    def text(c): return (c["Title"] + " " + c["Expected Result"]).lower()
    def groups(t, k): return sum(any(has(t, alt) for alt in grp.split("|")) for grp in k["Keywords"].split(";"))
    found = 0
    lines = []
    for k in keys:
        want = [x.strip() for x in k.get("Requirement", "").split("|") if x.strip()]
        pool = [c for c in cases if not want or set(refs(c["Requirements"])) & set(want)]
        n_groups = len(k["Keywords"].split(";"))
        hits = [c for c in pool if groups(text(c), k) == n_groups]
        words = {w for w in re.findall(r"[a-z0-9]{4,}", k["Condition"].lower())}
        hit = max(hits, key=lambda c: (groups(c["Title"].lower(), k), len(words & set(re.findall(r"[a-z0-9]{4,}", c["Title"].lower())))),
                  default=None)  # the case whose title says it best
        found += hit is not None
        lines.append(f"  {'FOUND ' if hit else 'MISSED'} {k['Key']} {k['Condition']}" + (f"\n         <- {hit['ID']}: {hit['Title'][:120]}" if hit else ""))
    n = len(keys)
    print(f"SCORE {module}: {found} of {n} required test conditions covered ({100 * found // max(n, 1)}%); {len(cases)} active cases")
    print("\n".join(lines))
    target = int(os.environ.get("TH_SCORE_TARGET", "80"))
    sys.exit(0 if 100 * found >= target * n else 1)


def cmd_baseline(a):
    if len(a) < 2: die("Usage: design.py baseline <MODULE> <approved-by>")
    module, by = a[0].upper(), a[1]
    cases = module_rows(read_csv(CASE_FILE, CASE_COLS), module)
    asks = [q for q in read_csv(ASK_FILE, ASK_COLS) if q["Module"].upper() == module]
    b = baseline(module)
    base = os.path.join(TC_DIR, "baseline", module)
    os.makedirs(base, exist_ok=True)
    nums = [int(d[1:]) for d in os.listdir(base) if re.fullmatch(r"v\d+", d)]
    v = f"v{max(nums, default=0) + 1}"
    out = os.path.join(base, v); os.makedirs(out)
    write_csv(os.path.join(out, "TestCases.csv"), CASE_COLS, cases)
    write_csv(os.path.join(out, "requests.csv"), ASK_COLS, asks)
    with open(os.path.join(out, "meta.json"), "w", encoding="utf-8", newline="\n") as f:
        json.dump({"module": module, "version": v, "approved_by": by, "approved_at": time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime()),
                   "analysis_baseline": b[0] if b else None, "cases": len(cases), "active": sum(1 for c in cases if active(c))},
                  f, indent=2, ensure_ascii=False)
    print(f"DESIGN BASELINE {module} {v}: {sum(1 for c in cases if active(c))} active cases (from analysis {b[0] if b else '?'}) -> {out.replace(os.sep, '/')}")


COMMANDS = {"check": cmd_check, "workbook": cmd_workbook, "import": cmd_import, "impact": cmd_impact, "score": cmd_score, "baseline": cmd_baseline}
if __name__ == "__main__":
    if len(sys.argv) < 2 or sys.argv[1] not in COMMANDS: die("Usage: design.py " + "|".join(COMMANDS) + " ...")
    COMMANDS[sys.argv[1]](sys.argv[2:])
