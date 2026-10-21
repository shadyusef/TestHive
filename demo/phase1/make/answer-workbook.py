"""Demo helper: plays the human. Fills the Gaps sheet of an analysis workbook from a CSV of answers.
Usage: python answer-workbook.py <workbook.xlsx> <answers.csv>   (answers.csv columns: ID, Answer, Answered By)"""
import sys, csv
import openpyxl

wb_path, answers_path = sys.argv[1], sys.argv[2]
with open(answers_path, encoding="utf-8-sig", newline="") as f:
    answers = {r["ID"]: r for r in csv.DictReader(f)}
wb = openpyxl.load_workbook(wb_path)
ws = wb["Gaps"]
head = [c.value for c in ws[1]]
col = {h: head.index(h) + 1 for h in head}
filled = 0
for row in range(2, ws.max_row + 1):
    a = answers.get(ws.cell(row, col["ID"]).value)
    if not a: continue
    ws.cell(row, col["Answer"]).value = a["Answer"]
    ws.cell(row, col["Answered By"]).value = a["Answered By"]
    ws.cell(row, col["Status"]).value = "Answered"
    filled += 1
wb.save(wb_path)
print(f"Filled {filled} answers in {wb_path} (as a human would in Excel)")
