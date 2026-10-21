"""Writes demo/phase1/sources/user-stories.docx (fictional Nakhla Rewards stories with planted gaps).
Usage: python demo/phase1/make/make-stories.py   (needs python-docx; run from the TestHive folder)"""
import docx

d = docx.Document()
d.add_heading("Nakhla Rewards: Sign-up user stories", 0)
d.add_paragraph("Sprint 14 backlog. Fictional demo document for TestHive phase 1.")
d.add_heading("Stories", 1)
t = d.add_table(rows=1, cols=3)
t.style = "Table Grid"
for c, h in zip(t.rows[0].cells, ("ID", "Story", "Acceptance criteria")): c.text = h
stories = [
    ("US-1", "As a visitor aged 16 or older, I can create an account so that I can collect points.",
     "Account is created after the mobile number is verified."),
    ("US-2", "As a visitor without a Saudi mobile number, I can sign up with my email address instead.",
     "A verification link is sent to the email."),
    ("US-3", "As a new member, I can request a new code if I did not receive the SMS.",
     "A new code is sent."),
    ("US-4", "As a new member, I receive my welcome points.",
     "The points balance shows 50."),
    ("US-5", "As a customer, I can choose the app language during sign-up.",
     "All screens appear in the chosen language."),
]
for s in stories:
    row = t.add_row().cells
    for c, v in zip(row, s): c.text = v
d.add_heading("Notes from refinement", 1)
d.add_paragraph("The OTP screen follows the design mock-up (six digits).")
d.add_paragraph("Error messages will be provided by the content team later.")
d.save("demo/phase1/sources/user-stories.docx")
print("Wrote demo/phase1/sources/user-stories.docx")
